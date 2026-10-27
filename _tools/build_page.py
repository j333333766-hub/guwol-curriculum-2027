# 2027 입학생 3개년 편성표(학교스포츠클럽 반영안) 엑셀 → 한 장짜리 HTML
import openpyxl, json, html, sys
SRC = sys.argv[1]; OUT = sys.argv[2]
ws = openpyxl.load_workbook(SRC, data_only=True)['2027학년도 입학생']
rows = [list(r) for r in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True)]
n = lambda v: v if isinstance(v, (int, float)) else 0
cls = rows[4][9]

subj = []           # 교과 (학교자율시간 포함)
grp = ''
for r in rows[7:27]:
    if r[0]: pass
    g = (r[1] or '').strip('/ ') if r[1] else ''
    if g in ('소계',): break
    if g: grp = g
    name = (r[2] or r[1] or '').strip('/ ')
    if not name: continue
    sem = [n(r[8]), n(r[11]), n(r[12]), n(r[13]), n(r[14]), n(r[15])]
    if not any(sem) and not n(r[17]): continue
    subj.append(dict(grp=grp if grp not in ('학교', '자율', '시간') else '학교자율시간', name=name, base=r[4], adj=n(r[7]), sem=sem, pure=n(r[16]), free=n(r[23]), total=n(r[17]),
                     freeparts=[n(r[19]), n(r[20]), n(r[21]), n(r[22])]))
# 교과 소계, 창체, 총계
def pick(label_col, label):
    for r in rows:
        if r[label_col] and str(r[label_col]).replace('\n', '') .startswith(label): return r
sub = rows[27]
cc = []
for r in rows[28:36]:
    lab = ' '.join(str(x).replace('\n', '') for x in r[1:4] if x)
    if not lab or lab.startswith('소계'): continue
    sem = [n(r[8]), n(r[11]), n(r[12]), n(r[13]), n(r[14]), n(r[15])]
    cc.append(dict(name=lab, sem=sem, total=n(r[17])))
ccsub = rows[36]; tot = rows[37]; cnt = rows[38]
sports = [c for c in cc if '스포츠' in c['name'] or c['name'] in ('창체 순증', '교과 감축', '체육대체')]
import os
VER = sys.argv[4] if len(sys.argv) > 4 else ''
data = dict(ver=VER, src=os.path.basename(SRC), cls=cls, subj=subj, cc=cc,
            sub=dict(sem=[n(sub[8]), n(sub[11]), n(sub[12]), n(sub[13]), n(sub[14]), n(sub[15])], base=sub[4], total=n(sub[17]), free=n(sub[23])),
            ccsub=dict(sem=[n(ccsub[8]), n(ccsub[11]), n(ccsub[12]), n(ccsub[13]), n(ccsub[14]), n(ccsub[15])], base=ccsub[4], total=n(ccsub[17])),
            tot=dict(sem=[n(tot[8]), n(tot[11]), n(tot[12]), n(tot[13]), n(tot[14]), n(tot[15])], base=tot[4], total=n(tot[17]), free=n(tot[23])),
            cnt=[cnt[6], cnt[9], cnt[12], cnt[13], cnt[14], cnt[15]])
print(json.dumps(data, ensure_ascii=False)[:1500])
tpl = open(sys.argv[3], encoding='utf-8').read()
open(OUT, 'w', encoding='utf-8').write(tpl.replace('/*DATA*/null', json.dumps(data, ensure_ascii=False)))
